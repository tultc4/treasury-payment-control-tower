from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import math
import re
import sys
import unicodedata
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

from dateutil import parser as date_parser
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    import pdfplumber  # type: ignore
except Exception:  # pragma: no cover - PDF support is optional at runtime.
    pdfplumber = None


APP_NAME = "Bank Statement Transaction Parser"
APP_VERSION = "1.0.0"
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".txt", ".pdf"}

# The first 16 fields intentionally match Bank_Transaction_Detail_Template.xlsx
# and Treasury Payment Control Tower v6's bank statement importer.
TRANSACTION_OUTPUT_COLUMNS = [
    "BANK_NAME",
    "ACCOUNT_NUMBER",
    "ACCOUNT_CURRENCY",
    "BANK_REFERENCE",
    "CUSTOMER_REFERENCE",
    "ENTRY_DATE",
    "TRANSACTION_AMOUNT",
    "DEBIT_CREDIT_INDICATOR",
    "PRODUCT_TYPE",
    "TRANSACTION_DESCRIPTION",
    "EXTRA_INFORMATION",
    "PAYMENT_DETAILS",
    "BENEFICIARY_REFERENCE",
    "REMITTER_REFERENCE",
    "SOURCE_FILE",
    "SOURCE_SHEET",
    # Additional audit and matching fields.
    "BANK_GROUP",
    "ACCOUNT_NAME",
    "STATEMENT_DATE",
    "STATEMENT_NO",
    "VALUE_DATE",
    "BENEFICIARY_NAME",
    "BENEFICIARY_ACCOUNT",
    "END_TO_END_REFERENCE",
    "BANK_TRANSACTION_ID",
    "EFORM_CANDIDATE",
    "ENTITY",
    "SOURCE_ROW",
    "SOURCE_PAGE",
    "RAW_TEXT",
    "RECORD_STATUS",
    "VALIDATION_MESSAGE",
    "RECORD_KEY",
]

BALANCE_OUTPUT_COLUMNS = [
    "source_file_name",
    "source_type",
    "source_bank",
    "bank_group",
    "bank_name",
    "account_name",
    "account_number",
    "currency",
    "statement_date",
    "statement_no",
    "opening_balance",
    "closing_balance",
    "entity_code",
    "entity_name_short",
    "entity_name",
    "region",
    "mapping_status",
    "mapping_key",
    "mapping_note",
    "page_no",
]

# NOTE (fixed): this used to require a literal "Source File Name" column,
# a field the parser itself generates for the OUTPUT — no real bank export
# ever prints that as a header, so balance-summary detection could never
# succeed on real input. detect_balance_summary_header() below now goes
# through the same alias system as transaction detection instead of this set.

# Field aliases are deliberately broad. Exact bank-specific mappings can be added
# in parser_config.json without changing this source file.
DEFAULT_HEADER_ALIASES: dict[str, list[str]] = {
    "BANK_NAME": [
        "Bank Name",
        "Bank",
        "Financial Institution",
    ],
    "BANK_GROUP": [
        "Bank Group",
        "Source Bank",
        "Bank Code",
    ],
    "ACCOUNT_NAME": [
        "Account Name",
        "Account Holder",
        "Account Holder Name",
        "Customer Name",
    ],
    "ACCOUNT_NUMBER": [
        "Account Number",
        "Account No",
        "Account No.",
        "A/C Number",
        "A/C No",
        "Acct Number",
        "Acct No",
        "Debit Account",
        "Ordering Account",
        "Account",
    ],
    "ACCOUNT_CURRENCY": [
        "Account Currency",
        "Currency",
        "CCY",
        "Acct Currency",
        "Account CCY",
    ],
    "STATEMENT_DATE": [
        "Statement Date",
        "As Of Date",
        "Report Date",
        "Closing Date",
    ],
    "STATEMENT_NO": [
        "Statement Number",
        "Statement No",
        "Statement No.",
    ],
    "ENTRY_DATE": [
        "Entry Date",
        "Posting Date",
        "Booked Date",
        "Booking Date",
        "Book Date",
        "Transaction Date",
        "Txn Date",
        "Date",
    ],
    "VALUE_DATE": [
        "Value Date",
        "Val Date",
        "Effective Date",
    ],
    "TRANSACTION_AMOUNT": [
        "Transaction Amount",
        "Txn Amount",
        "Signed Amount",
        "Amount",
        "Local Amount",
        "Booking Amount",
    ],
    "DEBIT_AMOUNT": [
        "Debit Amount",
        "Debit",
        "Withdrawal",
        "Withdrawals",
        "Money Out",
        "Paid Out",
        "Debit Value",
    ],
    "CREDIT_AMOUNT": [
        "Credit Amount",
        "Credit",
        "Deposit",
        "Deposits",
        "Money In",
        "Paid In",
        "Credit Value",
    ],
    "DEBIT_CREDIT_INDICATOR": [
        "Debit Credit Indicator",
        "Debit/Credit Indicator",
        "D/C",
        "Dr Cr",
        "DR/CR",
        "Credit Debit",
        "Transaction Type Indicator",
        "Sign",
    ],
    "BANK_REFERENCE": [
        "Bank Reference",
        "Bank Ref",
        "Bank Reference Number",
        "Transaction Reference",
        "Transaction Ref",
        "Reference Number",
        "Reference No",
        "Reference",
    ],
    "CUSTOMER_REFERENCE": [
        "Customer Reference",
        "Customer Ref",
        "Client Reference",
        "Client Ref",
        "Your Reference",
        "Payment Reference",
        "Payment Ref",
        "Eform Number",
        "E-form Number",
        "Eform",
        "E-form",
    ],
    "PRODUCT_TYPE": [
        "Product Type",
        "Transaction Type",
        "Payment Type",
        "Transaction Code",
        "Txn Code",
    ],
    "TRANSACTION_DESCRIPTION": [
        "Transaction Description",
        "Description",
        "Transaction Details",
        "Narrative",
        "Particulars",
        "Details",
        "Remarks",
        "Memo",
    ],
    "EXTRA_INFORMATION": [
        "Extra Information",
        "Additional Information",
        "Additional Details",
        "Supplementary Details",
    ],
    "PAYMENT_DETAILS": [
        "Payment Details",
        "Remittance Information",
        "Remittance Details",
        "Payment Information",
        "Payment Narrative",
    ],
    "BENEFICIARY_REFERENCE": [
        "Beneficiary Reference",
        "Beneficiary Ref",
        "Payee Reference",
    ],
    "REMITTER_REFERENCE": [
        "Remitter Reference",
        "Remitter Ref",
        "Payer Reference",
    ],
    "BENEFICIARY_NAME": [
        "Beneficiary Name",
        "Payee Name",
        "Counterparty Name",
    ],
    "BENEFICIARY_ACCOUNT": [
        "Beneficiary Account",
        "Beneficiary Account Number",
        "Payee Account",
        "Counterparty Account",
    ],
    "END_TO_END_REFERENCE": [
        "End To End Reference",
        "End-to-End Reference",
        "End To End ID",
        "End-to-End ID",
    ],
    "BANK_TRANSACTION_ID": [
        "Bank Transaction ID",
        "Transaction ID",
        "Txn ID",
        "Bank ID",
    ],
    "OPENING_BALANCE": [
        "Opening Balance",
        "Beginning Balance",
        "Balance Brought Forward",
        "Opening Ledger Balance",
        "Opening Available Balance",
    ],
    "CLOSING_BALANCE": [
        "Closing Balance",
        "Ending Balance",
        "Balance Carried Forward",
        "Current / Closing Ledger Balance",
        "Closing Ledger Balance",
        "Current Ledger Balance",
        "Ledger Balance",
    ],
    # A per-transaction running balance ("Balance" on Standard Chartered's
    # export). Never used to detect a balance-summary table — only read by
    # derive_balances_from_transaction_table(), where the LAST row's value
    # per account is that account's closing balance.
    "RUNNING_BALANCE": [
        "Running Balance",
        "Balance",
        "Book Balance",
        "Balance After Transaction",
    ],
}

DEFAULT_METADATA_ALIASES: dict[str, list[str]] = {
    "BANK_NAME": ["Bank Name", "Bank"],
    "ACCOUNT_NAME": ["Account Name", "Account Holder", "Customer Name", "Account Nickname"],
    "ACCOUNT_NUMBER": ["Account Number", "Account No", "A/C No", "Acct No"],
    "ACCOUNT_CURRENCY": ["Account Currency", "Currency", "CCY"],
    "STATEMENT_DATE": ["Statement Date", "As Of Date", "Report Date"],
    "STATEMENT_NO": ["Statement Number", "Statement No"],
    # "Ledger Balance" is the current/running book balance a bank actually
    # prints on simple exports (no separate opening/closing pair) — treat it
    # as the closing balance. "Available Balance" (funds minus holds) is a
    # lower-priority fallback for the same purpose, used only when nothing
    # better is found (see extract_balance_from_metadata()).
    "OPENING_BALANCE": ["Opening Balance", "Beginning Balance", "Opening Ledger Balance"],
    "CLOSING_BALANCE": ["Closing Balance", "Ending Balance", "Ledger Balance", "Current / Closing Ledger Balance"],
    "AVAILABLE_BALANCE": ["Available Balance", "Available Ledger Balance"],
}

DEFAULT_CONFIG: dict[str, Any] = {
    "header_aliases": DEFAULT_HEADER_ALIASES,
    "metadata_aliases": DEFAULT_METADATA_ALIASES,
    "bank_patterns": [
        {
            # NOTE: underscore is a \w character, so \bCITI\b would not match
            # a filename like "Citi_01_3107_gamco.xlsx" (no boundary between
            # "Citi" and "_"). Use the same safe non-alnum boundary as the
            # DB/HSBC/SC patterns below instead of \b.
            "pattern": r"(?:^|[^A-Z0-9])CITI(?:BANK)?(?:[^A-Z0-9]|$)",
            "bank_group": "CITI",
            "bank_name": "CITIBANK",
        },
        {
            "pattern": r"\bDEUTSCHE\b|(?:^|[^A-Z0-9])DB(?:[^A-Z0-9]|$)",
            "bank_group": "DB",
            "bank_name": "DEUTSCHE BANK SINGAPORE BRANCH",
        },
        {
            # Same underscore-boundary issue as CITI above (e.g.
            # "HSBC_01_3107_001.xlsx") — use the safe non-alnum boundary.
            "pattern": r"(?:^|[^A-Z0-9])HSBC(?:[^A-Z0-9]|$)|HONGKONG AND SHANGHAI|HK AND SHANGHAI",
            "bank_group": "HSBC",
            "bank_name": "HSBC",
        },
        {
            "pattern": r"STANDARD CHARTERED|STANCHART|(?:^|[^A-Z0-9])SC(?:[^A-Z0-9]|$)",
            "bank_group": "SC",
            "bank_name": "STANDARD CHARTERED",
        },
        {
            "pattern": r"BANGKOK\s*BANK|(?:^|[^A-Z0-9])BBL(?:[^A-Z0-9]|$)|(?:^|[^A-Z0-9])BK(?:[^A-Z0-9]|$)",
            "bank_group": "BANGKOK",
            "bank_name": "BANGKOK BANK",
        },
    ],
    "skip_description_patterns": [
        r"\bOPENING BALANCE\b",
        r"\bCLOSING BALANCE\b",
        r"\bBEGINNING BALANCE\b",
        r"\bENDING BALANCE\b",
        r"\bBALANCE BROUGHT FORWARD\b",
        r"\bBALANCE CARRIED FORWARD\b",
        r"\bBROUGHT FORWARD\b",
        r"\bCARRIED FORWARD\b",
        r"\bTOTAL DEBIT\b",
        r"\bTOTAL CREDIT\b",
        r"\bSUBTOTAL\b",
    ],
    "eform_patterns": [
        r"\b[A-Z]{1,8}-PM\d{6,14}(?:-\d+)?\b",
        r"\bPM\d{8,16}(?:-\d+)?\b",
        r"\bFA[-_/ ]?PM\d{6,14}(?:-\d+)?\b",
    ],
    "date_dayfirst_default": True,
    "date_dayfirst_by_bank": {
        "CITI": True,
        "DB": True,
        "HSBC": True,
        "SC": True,
    },
    "scan_header_rows": 80,
    "scan_metadata_rows": 60,
    "scan_metadata_columns": 40,
    "minimum_header_score": 8,
    "forward_fill_dates": True,
    "use_statement_date_when_entry_date_missing": False,
    "exclude_zero_amount": True,
    "account_master_file": "account_master.csv",
    "account_overrides": {},
}

NAVY = "17365D"
DARK_BLUE = "1F4E78"
TEAL = "008C95"
LIGHT_BLUE = "DDEBF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_RED = "F4CCCC"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "7F7F7F"
RED = "C00000"
GREEN = "008000"
BLUE = "0000FF"
ORANGE = "ED7D31"

AMOUNT_FORMAT = '#,##0.00;[Red](#,##0.00);-'
DATE_FORMAT = "dd-mmm-yy"
DATETIME_FORMAT = "dd-mmm-yy hh:mm"


@dataclass
class ParseLogEntry:
    source_file: str
    source_type: str
    status: str
    bank_group: str = ""
    bank_name: str = ""
    sheets_or_pages_scanned: int = 0
    transaction_rows: int = 0
    balance_rows: int = 0
    skipped_rows: int = 0
    review_rows: int = 0
    duplicate_rows: int = 0
    message: str = ""


@dataclass
class ParseResult:
    transactions: list[dict[str, Any]] = field(default_factory=list)
    balances: list[dict[str, Any]] = field(default_factory=list)
    logs: list[ParseLogEntry] = field(default_factory=list)
    unmapped_rows: list[dict[str, Any]] = field(default_factory=list)
    duplicates: list[dict[str, Any]] = field(default_factory=list)


class ParserError(RuntimeError):
    pass


_ILLEGAL_XML_CHARS_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def xml_safe(value: Any) -> Any:
    """Strips ASCII control characters openpyxl/Excel XML rejects outright
    (IllegalCharacterError). Seen in practice from PDF table extraction
    (pdfplumber occasionally emits stray control bytes between letters,
    e.g. a spaced-out " C I T I B A N K " bank name). Only strings are
    affected; numbers/dates pass through untouched."""
    if isinstance(value, str):
        return _ILLEGAL_XML_CHARS_RE.sub("", value)
    return value


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    text = xml_safe(str(value)).replace("\xa0", " ").replace("\u200b", " ")
    return re.sub(r"\s+", " ", text).strip()


def ascii_fold(value: Any) -> str:
    text = clean_text(value)
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )


def norm_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", ascii_fold(value).upper())


def norm_ref(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", ascii_fold(value).upper())


def account_key(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    # Excel may turn a text account number into a numeric value with .0.
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    digits = re.sub(r"\D", "", text)
    if digits:
        return digits.lstrip("0") or "0"
    return norm_ref(text)


def display_account(value: Any) -> str:
    text = clean_text(value)
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    # Keep only common visual separators; dashboard matching normalizes separately.
    return text.strip()


def normalize_currency(value: Any) -> str:
    text = ascii_fold(value).upper().strip()
    match = re.search(r"\b([A-Z]{3})\b", text)
    return match.group(1) if match else text[:3]


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)

    text = clean_text(value)
    if not text:
        return None

    upper = text.upper()
    negative = False
    if upper.startswith("(") and upper.endswith(")"):
        negative = True
    if re.search(r"(?:^|\s)(DR|DEBIT)(?:\s|$)", upper) or upper.endswith("DR"):
        negative = True
    if re.search(r"(?:^|\s)(CR|CREDIT)(?:\s|$)", upper) or upper.endswith("CR"):
        negative = False
    if upper.startswith("-"):
        negative = True

    # Retain digits and decimal/thousand separators only.
    raw = re.sub(r"[^0-9,\.\-+]", "", text)
    if not re.search(r"\d", raw):
        return None

    sign = -1.0 if negative else 1.0
    if raw.startswith("-"):
        sign = -1.0
    raw = raw.lstrip("+-")

    if "," in raw and "." in raw:
        # The final separator is treated as the decimal separator.
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        # A single comma followed by one or two digits is most likely decimal.
        parts = raw.split(",")
        if len(parts) == 2 and 1 <= len(parts[1]) <= 2:
            raw = parts[0] + "." + parts[1]
        else:
            raw = raw.replace(",", "")
    elif raw.count(".") > 1:
        parts = raw.split(".")
        if len(parts[-1]) in {1, 2}:
            raw = "".join(parts[:-1]) + "." + parts[-1]
        else:
            raw = "".join(parts)

    try:
        return sign * abs(float(raw))
    except ValueError:
        return None


def parse_date_value(value: Any, dayfirst: bool = True) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # A raw Excel serial may appear in CSV exports or malformed workbooks.
        if 1 <= float(value) <= 100000:
            try:
                base = datetime(1899, 12, 30)
                return (base + timedelta(days=float(value))).date()
            except Exception:
                pass
    text = clean_text(value)
    if not text:
        return None
    try:
        if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", text):
            parsed = date_parser.parse(text, yearfirst=True, dayfirst=False, fuzzy=False)
        else:
            parsed = date_parser.parse(text, dayfirst=dayfirst, fuzzy=False)
        return parsed.date()
    except (ValueError, OverflowError, TypeError):
        return None


def to_excel_datetime(value: date | datetime | None) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    return datetime.combine(value, datetime.min.time())


def indicator_value(value: Any) -> str:
    text = ascii_fold(value).upper().strip()
    if not text:
        return ""
    if text in {"D", "DR", "DEBIT", "DB", "WITHDRAWAL", "OUT"}:
        return "DEBIT"
    if text in {"C", "CR", "CREDIT", "CRDT", "DEPOSIT", "IN"}:
        return "CREDIT"
    if "DEBIT" in text or re.search(r"\bDR\b", text):
        return "DEBIT"
    if "CREDIT" in text or re.search(r"\bCR\b", text):
        return "CREDIT"
    return text


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in base.items():
        out[key] = copy(value)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = value
    return out


def load_config(path: Path | None) -> dict[str, Any]:
    config = DEFAULT_CONFIG
    if path:
        if not path.exists():
            raise FileNotFoundError(f"Config file not found: {path}")
        with path.open("r", encoding="utf-8-sig") as fh:
            user_config = json.load(fh)
        config = deep_merge(DEFAULT_CONFIG, user_config)
    return config


def compile_alias_index(config: dict[str, Any]) -> dict[str, str]:
    index: dict[str, str] = {}
    for canonical, aliases in config["header_aliases"].items():
        for alias in [canonical, *aliases]:
            key = norm_header(alias)
            if key and key not in index:
                index[key] = canonical
    return index


def compile_metadata_index(config: dict[str, Any]) -> dict[str, str]:
    index: dict[str, str] = {}
    for canonical, aliases in config["metadata_aliases"].items():
        for alias in [canonical, *aliases]:
            key = norm_header(alias)
            if key and key not in index:
                index[key] = canonical
    return index


def load_account_master(path: Path | None) -> tuple[dict[tuple[str, str, str], dict[str, str]], dict[str, dict[str, str]]]:
    exact: dict[tuple[str, str, str], dict[str, str]] = {}
    by_alias: dict[str, dict[str, str]] = {}
    if not path or not path.exists():
        return exact, by_alias

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            bank = clean_text(row.get("BANK_GROUP")).upper()
            currency = normalize_currency(row.get("CURRENCY"))
            normalized = display_account(row.get("NORMALIZED_ACCOUNT"))
            if not normalized:
                continue
            data = {k: clean_text(v) for k, v in row.items() if k}
            exact[(bank, account_key(normalized), currency)] = data
            aliases = clean_text(row.get("RAW_ACCOUNT_ALIASES"))
            for alias in re.split(r"[|;,]+", aliases):
                alias_key = account_key(alias)
                if alias_key:
                    by_alias[f"{bank}|{alias_key}|{currency}"] = data
            by_alias[f"{bank}|{account_key(normalized)}|{currency}"] = data
    return exact, by_alias


def resolve_account(
    bank_group: str,
    raw_account: Any,
    currency: str,
    account_master: tuple[dict[tuple[str, str, str], dict[str, str]], dict[str, dict[str, str]]],
    config: dict[str, Any],
) -> tuple[str, str, str]:
    """Return display account, entity, mapping note."""
    raw_display = display_account(raw_account)
    raw_key = account_key(raw_account)
    bank = clean_text(bank_group).upper()
    ccy = normalize_currency(currency)

    overrides = config.get("account_overrides", {})
    override_key = f"{bank}|{raw_key}|{ccy}"
    if override_key in overrides:
        raw_display = clean_text(overrides[override_key])
        raw_key = account_key(raw_display)

    exact, aliases = account_master
    data = exact.get((bank, raw_key, ccy)) or aliases.get(override_key) or aliases.get(f"{bank}|{raw_key}|{ccy}")
    if data:
        account = clean_text(data.get("NORMALIZED_ACCOUNT")) or raw_display
        entity = clean_text(data.get("ENTITY"))
        return account, entity, "ACCOUNT_MASTER"
    return raw_display, "", "UNMAPPED_ACCOUNT" if raw_display else "MISSING_ACCOUNT"


def resolve_account_currency_hint(
    bank_group: str,
    raw_account: Any,
    account_master: tuple[dict[tuple[str, str, str], dict[str, str]], dict[str, dict[str, str]]],
) -> str:
    """Some simple balance exports (e.g. a plain "Account No. / Ledger
    Balance" snippet) never print a currency at all. account_master.csv
    already knows the currency for every account we track, so look it up
    there instead of leaving the balance row uncategorized."""
    bank = clean_text(bank_group).upper()
    raw_key = account_key(raw_account)
    if not raw_key:
        return ""
    _, aliases = account_master
    prefix = f"{bank}|{raw_key}|"
    for key, data in aliases.items():
        if key.startswith(prefix):
            return clean_text(data.get("CURRENCY"))
    return ""


def infer_bank(text: str, config: dict[str, Any]) -> tuple[str, str]:
    upper = ascii_fold(text).upper()
    for item in config.get("bank_patterns", []):
        if re.search(item["pattern"], upper, flags=re.IGNORECASE):
            return clean_text(item.get("bank_group")).upper(), clean_text(item.get("bank_name"))
    return "", ""


def detect_eform(values: Iterable[Any], config: dict[str, Any]) -> str:
    text = " | ".join(clean_text(v) for v in values if clean_text(v))
    upper = ascii_fold(text).upper()
    candidates: list[str] = []
    for pattern in config.get("eform_patterns", []):
        for match in re.finditer(pattern, upper, flags=re.IGNORECASE):
            candidate = re.sub(r"[\s_/]+", "-", match.group(0).upper())
            if candidate not in candidates:
                candidates.append(candidate)
    # Prefer the most specific/longest reference and suppress substring duplicates,
    # e.g. keep FA-PM260722001 instead of also returning PM260722001.
    kept: list[str] = []
    for candidate in sorted(candidates, key=len, reverse=True):
        if any(candidate in existing or existing in candidate for existing in kept):
            continue
        kept.append(candidate)
    return " | ".join(kept)


def record_key(record: dict[str, Any]) -> str:
    txid = norm_ref(record.get("BANK_TRANSACTION_ID"))
    if txid:
        core = [clean_text(record.get("BANK_GROUP")).upper(), account_key(record.get("ACCOUNT_NUMBER")), txid]
    else:
        entry = parse_date_value(record.get("ENTRY_DATE"))
        amount = parse_number(record.get("TRANSACTION_AMOUNT")) or 0.0
        core = [
            clean_text(record.get("BANK_GROUP")).upper(),
            account_key(record.get("ACCOUNT_NUMBER")),
            normalize_currency(record.get("ACCOUNT_CURRENCY")),
            entry.isoformat() if entry else "",
            f"{amount:.6f}",
            norm_ref(record.get("BANK_REFERENCE")),
            norm_ref(record.get("CUSTOMER_REFERENCE")),
            norm_ref(record.get("TRANSACTION_DESCRIPTION"))[:120],
        ]
    return hashlib.sha256("|".join(core).encode("utf-8")).hexdigest()[:24].upper()


def row_raw_text(values: Sequence[Any]) -> str:
    return " | ".join(clean_text(v) for v in values if clean_text(v))


def row_looks_like_balance_or_total(record: dict[str, Any], config: dict[str, Any]) -> bool:
    text = " ".join(
        clean_text(record.get(field))
        for field in [
            "TRANSACTION_DESCRIPTION",
            "PRODUCT_TYPE",
            "PAYMENT_DETAILS",
            "EXTRA_INFORMATION",
            "RAW_TEXT",
        ]
    ).upper()
    return any(re.search(pattern, text) for pattern in config.get("skip_description_patterns", []))


def dayfirst_for_bank(bank_group: str, config: dict[str, Any]) -> bool:
    per_bank = config.get("date_dayfirst_by_bank", {})
    if bank_group in per_bank:
        return bool(per_bank[bank_group])
    return bool(config.get("date_dayfirst_default", True))


def match_header_cell(value: Any, alias_index: dict[str, str]) -> str | None:
    key = norm_header(value)
    if not key:
        return None
    if key in alias_index:
        return alias_index[key]
    # A bank export may append a currency or explanatory suffix to a header.
    candidates: list[tuple[int, str]] = []
    for alias, canonical in alias_index.items():
        if len(alias) >= 5 and (alias in key or key in alias):
            candidates.append((len(alias), canonical))
    if not candidates:
        return None
    candidates.sort(reverse=True)
    return candidates[0][1]


def effective_data_width(rows: Sequence[Sequence[Any]], header_idx: int, sample: int = 10) -> int:
    """Some real exports append a long field-name "catalog" far to the right
    of the actual header row (seen on real Citi and HSBC files: dozens of
    extra column labels like "Credit Amount"/"Debit Amount"/"Value Date"
    that describe fields the export COULD carry, with nothing underneath
    them in the real data rows). Left unchecked, those phantom labels get
    picked up by the alias matcher and corrupt header detection — e.g. a
    pure balance-only file gets misread as having a transaction-amount
    column. Cap scanning to the widest column that actually has data in a
    sample of the rows below the header, not just wherever the header row
    happens to have text."""
    widest = 0
    for row in rows[header_idx + 1 : header_idx + 1 + sample]:
        for idx in range(len(row) - 1, -1, -1):
            if clean_text(row[idx]):
                widest = max(widest, idx + 1)
                break
    return widest


def header_mapping_from_rows(
    row1: Sequence[Any],
    row2: Sequence[Any] | None,
    alias_index: dict[str, str],
) -> dict[str, int]:
    width = max(len(row1), len(row2) if row2 else 0)
    mapping: dict[str, int] = {}
    for idx in range(width):
        value1 = row1[idx] if idx < len(row1) else None
        value2 = row2[idx] if row2 and idx < len(row2) else None
        choices = [value1, value2, f"{clean_text(value1)} {clean_text(value2)}"]
        canonical = None
        for choice in choices:
            canonical = match_header_cell(choice, alias_index)
            if canonical:
                break
        if canonical and canonical not in mapping:
            mapping[canonical] = idx
    return mapping


def header_score(mapping: dict[str, int]) -> int:
    score = 0
    if "ENTRY_DATE" in mapping:
        score += 4
    if "TRANSACTION_AMOUNT" in mapping:
        score += 5
    if "DEBIT_AMOUNT" in mapping or "CREDIT_AMOUNT" in mapping:
        score += 4
    if "ACCOUNT_NUMBER" in mapping:
        score += 3
    if "ACCOUNT_CURRENCY" in mapping:
        score += 2
    if "TRANSACTION_DESCRIPTION" in mapping:
        score += 2
    if "BANK_REFERENCE" in mapping or "CUSTOMER_REFERENCE" in mapping:
        score += 2
    if "VALUE_DATE" in mapping:
        score += 1
    return score


def detect_transaction_header(
    rows: Sequence[Sequence[Any]],
    config: dict[str, Any],
) -> tuple[int, int, dict[str, int]] | None:
    alias_index = compile_alias_index(config)
    best: tuple[int, int, dict[str, int], int] | None = None
    max_rows = min(len(rows), int(config.get("scan_header_rows", 80)))
    for idx in range(max_rows):
        width = effective_data_width(rows, idx) or len(rows[idx])
        row1 = rows[idx][:width]
        mapping1 = header_mapping_from_rows(row1, None, alias_index)
        score1 = header_score(mapping1)
        if best is None or score1 > best[3]:
            best = (idx, idx, mapping1, score1)

        if idx + 1 < max_rows:
            mapping2 = header_mapping_from_rows(row1, rows[idx + 1][:width], alias_index)
            score2 = header_score(mapping2)
            if best is None or score2 > best[3]:
                best = (idx, idx + 1, mapping2, score2)

    if not best:
        return None
    start, end, mapping, score = best
    has_amount = "TRANSACTION_AMOUNT" in mapping or "DEBIT_AMOUNT" in mapping or "CREDIT_AMOUNT" in mapping
    has_date = "ENTRY_DATE" in mapping or "VALUE_DATE" in mapping
    if score < int(config.get("minimum_header_score", 8)) or not has_amount or not has_date:
        return None
    return start, end, mapping


def detect_balance_summary_header(
    rows: Sequence[Sequence[Any]],
    config: dict[str, Any],
) -> tuple[int, dict[str, int]] | None:
    """Finds a table of one-row-per-account-per-day balances (e.g. a Citi
    "Opening/Closing Ledger Balance" export). Uses the same alias index as
    detect_transaction_header() so real header wording ("Ledger Balance",
    "Account Currency", ...) resolves correctly. A row that also has a real
    transaction-amount column is a transaction table, not a balance summary,
    even if it happens to carry a running balance alongside — leave that to
    detect_transaction_header()."""
    alias_index = compile_alias_index(config)
    max_rows = min(len(rows), int(config.get("scan_header_rows", 80)))
    for idx in range(max_rows):
        width = effective_data_width(rows, idx) or len(rows[idx])
        mapping = header_mapping_from_rows(rows[idx][:width], None, alias_index)
        has_balance = "OPENING_BALANCE" in mapping or "CLOSING_BALANCE" in mapping
        has_account = "ACCOUNT_NUMBER" in mapping
        has_txn_amount = "TRANSACTION_AMOUNT" in mapping or "DEBIT_AMOUNT" in mapping or "CREDIT_AMOUNT" in mapping
        if has_balance and has_account and not has_txn_amount:
            return idx, mapping
    return None


def extract_metadata(rows: Sequence[Sequence[Any]], source_name: str, config: dict[str, Any]) -> dict[str, Any]:
    metadata_index = compile_metadata_index(config)
    metadata: dict[str, Any] = {}
    max_rows = min(len(rows), int(config.get("scan_metadata_rows", 60)))
    max_cols = int(config.get("scan_metadata_columns", 40))

    for r in range(max_rows):
        row = list(rows[r])[:max_cols]
        for c, value in enumerate(row):
            text = clean_text(value)
            if not text:
                continue
            key = norm_header(text)
            canonical = metadata_index.get(key)
            if canonical:
                for adjacent in row[c + 1 : c + 4]:
                    if clean_text(adjacent):
                        metadata.setdefault(canonical, adjacent)
                        break
            # Handle labels and values in the same cell, e.g. "Account No: 123".
            for label_key, field in metadata_index.items():
                label_text = re.sub(r"([A-Z])", r"\1", label_key)
                if label_key and key.startswith(label_key) and key != label_key:
                    value_part = re.split(r"[:=]", text, maxsplit=1)
                    if len(value_part) == 2 and clean_text(value_part[1]):
                        metadata.setdefault(field, value_part[1])

    full_text = source_name + "\n" + "\n".join(row_raw_text(row) for row in rows[:max_rows])
    bank_group, bank_name = infer_bank(full_text, config)
    metadata.setdefault("BANK_GROUP", bank_group)
    metadata.setdefault("BANK_NAME", bank_name)

    # Regex fallbacks for metadata that is often printed in a title block.
    if not clean_text(metadata.get("ACCOUNT_NUMBER")):
        match = re.search(
            r"(?:ACCOUNT\s*(?:NUMBER|NO\.?|#)|A/C\s*(?:NUMBER|NO\.?))\s*[:#-]?\s*([A-Z0-9][A-Z0-9\- ]{4,30})",
            ascii_fold(full_text).upper(),
        )
        if match:
            metadata["ACCOUNT_NUMBER"] = clean_text(match.group(1)).split("  ")[0]
    if not clean_text(metadata.get("ACCOUNT_CURRENCY")):
        match = re.search(r"(?:ACCOUNT\s*)?(?:CURRENCY|CCY)\s*[:#-]?\s*([A-Z]{3})\b", ascii_fold(full_text).upper())
        if match:
            metadata["ACCOUNT_CURRENCY"] = match.group(1)
    return metadata


def amount_from_row(getter, raw_values: Sequence[Any]) -> tuple[float | None, str, str]:
    """Return signed amount, indicator, note."""
    direct_raw = getter("TRANSACTION_AMOUNT")
    debit_raw = getter("DEBIT_AMOUNT")
    credit_raw = getter("CREDIT_AMOUNT")
    indicator = indicator_value(getter("DEBIT_CREDIT_INDICATOR"))

    direct = parse_number(direct_raw)
    debit = parse_number(debit_raw)
    credit = parse_number(credit_raw)

    # A single row genuinely having BOTH a unified "Transaction Amount"
    # column AND separate Debit/Credit Amount columns populated is a red
    # flag, not a legitimate export shape — real bank formats use one style
    # or the other for a given transaction. It happens on wide exports that
    # append a field-name "catalog" far to the right of the real header
    # (e.g. "Total Credit Amount"/"Total Debit Amount" labels with no real
    # per-row data under them, picked up as if they were real columns) —
    # confirmed on a real Citi statement where this silently replaced a
    # correct -83,721.10 transaction amount with -1 (a phantom "count"
    # column). Prefer the direct column whenever it's present; only fall
    # back to separate debit/credit columns when there is no direct amount
    # at all (genuinely debit/credit-column-only bank formats).
    if direct is None and (debit is not None or credit is not None):
        debit_abs = abs(debit or 0.0)
        credit_abs = abs(credit or 0.0)
        if debit_abs and credit_abs:
            return credit_abs - debit_abs, "MIXED", "BOTH_DEBIT_AND_CREDIT_COLUMNS_POPULATED"
        if debit_abs:
            return -debit_abs, "DEBIT", "SEPARATE_DEBIT_COLUMN"
        if credit_abs:
            return credit_abs, "CREDIT", "SEPARATE_CREDIT_COLUMN"

    if direct is None:
        return None, indicator, "MISSING_AMOUNT"

    if indicator == "DEBIT":
        return -abs(direct), "DEBIT", "AMOUNT_PLUS_INDICATOR"
    if indicator == "CREDIT":
        return abs(direct), "CREDIT", "AMOUNT_PLUS_INDICATOR"

    raw_text = clean_text(direct_raw).upper()
    if direct < 0 or " DR" in f" {raw_text}" or raw_text.endswith("DR") or "DEBIT" in raw_text:
        return -abs(direct), "DEBIT", "SIGNED_AMOUNT"
    if direct > 0 and (" CR" in f" {raw_text}" or raw_text.endswith("CR") or "CREDIT" in raw_text):
        return abs(direct), "CREDIT", "SIGNED_AMOUNT"
    return direct, "DEBIT" if direct < 0 else "CREDIT", "SIGNED_AMOUNT_ASSUMED"


def get_value(values: Sequence[Any], mapping: dict[str, int], field: str) -> Any:
    idx = mapping.get(field)
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def build_transaction_record(
    values: Sequence[Any],
    mapping: dict[str, int],
    metadata: dict[str, Any],
    source_file: str,
    source_sheet: str,
    source_row: int | str,
    source_page: int | None,
    previous_entry_date: date | None,
    account_master,
    config: dict[str, Any],
) -> tuple[dict[str, Any] | None, date | None, str]:
    def get(field: str) -> Any:
        value = get_value(values, mapping, field)
        if value in (None, ""):
            value = metadata.get(field)
        return value

    bank_group = clean_text(get("BANK_GROUP")).upper()
    bank_name = clean_text(get("BANK_NAME"))
    if not bank_group or not bank_name:
        inferred_group, inferred_name = infer_bank(f"{source_file} {source_sheet} {row_raw_text(values)}", config)
        bank_group = bank_group or inferred_group
        bank_name = bank_name or inferred_name

    dayfirst = dayfirst_for_bank(bank_group, config)
    entry_date = parse_date_value(get_value(values, mapping, "ENTRY_DATE"), dayfirst=dayfirst)
    value_date = parse_date_value(get_value(values, mapping, "VALUE_DATE"), dayfirst=dayfirst)
    statement_date = parse_date_value(metadata.get("STATEMENT_DATE"), dayfirst=dayfirst)
    if not entry_date and config.get("forward_fill_dates", True):
        entry_date = previous_entry_date
    if not entry_date and value_date:
        entry_date = value_date
    if not entry_date and config.get("use_statement_date_when_entry_date_missing", False):
        entry_date = statement_date

    amount, indicator, amount_note = amount_from_row(get, values)
    if amount is None:
        return None, previous_entry_date, "NO_AMOUNT"
    if config.get("exclude_zero_amount", True) and abs(amount) < 1e-12:
        return None, entry_date or previous_entry_date, "ZERO_AMOUNT"

    currency = normalize_currency(get("ACCOUNT_CURRENCY"))
    account, entity, account_mapping_note = resolve_account(
        bank_group,
        get("ACCOUNT_NUMBER"),
        currency,
        account_master,
        config,
    )

    record: dict[str, Any] = {
        "BANK_NAME": bank_name,
        "ACCOUNT_NUMBER": account,
        "ACCOUNT_CURRENCY": currency,
        "BANK_REFERENCE": clean_text(get("BANK_REFERENCE")),
        "CUSTOMER_REFERENCE": clean_text(get("CUSTOMER_REFERENCE")),
        "ENTRY_DATE": to_excel_datetime(entry_date),
        "TRANSACTION_AMOUNT": amount,
        "DEBIT_CREDIT_INDICATOR": indicator,
        "PRODUCT_TYPE": clean_text(get("PRODUCT_TYPE")),
        "TRANSACTION_DESCRIPTION": clean_text(get("TRANSACTION_DESCRIPTION")),
        "EXTRA_INFORMATION": clean_text(get("EXTRA_INFORMATION")),
        "PAYMENT_DETAILS": clean_text(get("PAYMENT_DETAILS")),
        "BENEFICIARY_REFERENCE": clean_text(get("BENEFICIARY_REFERENCE")),
        "REMITTER_REFERENCE": clean_text(get("REMITTER_REFERENCE")),
        "SOURCE_FILE": source_file,
        "SOURCE_SHEET": source_sheet,
        "BANK_GROUP": bank_group,
        "ACCOUNT_NAME": clean_text(get("ACCOUNT_NAME")),
        "STATEMENT_DATE": to_excel_datetime(statement_date),
        "STATEMENT_NO": clean_text(get("STATEMENT_NO")),
        "VALUE_DATE": to_excel_datetime(value_date),
        "BENEFICIARY_NAME": clean_text(get("BENEFICIARY_NAME")),
        "BENEFICIARY_ACCOUNT": display_account(get("BENEFICIARY_ACCOUNT")),
        "END_TO_END_REFERENCE": clean_text(get("END_TO_END_REFERENCE")),
        "BANK_TRANSACTION_ID": clean_text(get("BANK_TRANSACTION_ID")),
        "ENTITY": entity,
        "SOURCE_ROW": source_row,
        "SOURCE_PAGE": source_page,
        "RAW_TEXT": row_raw_text(values),
    }

    record["EFORM_CANDIDATE"] = detect_eform(
        [
            record["CUSTOMER_REFERENCE"],
            record["BANK_REFERENCE"],
            record["END_TO_END_REFERENCE"],
            record["PAYMENT_DETAILS"],
            record["TRANSACTION_DESCRIPTION"],
            record["EXTRA_INFORMATION"],
            record["BENEFICIARY_REFERENCE"],
            record["REMITTER_REFERENCE"],
        ],
        config,
    )

    messages: list[str] = []
    if not bank_group:
        messages.append("MISSING_BANK_GROUP")
    if not account:
        messages.append("MISSING_ACCOUNT_NUMBER")
    if not currency:
        messages.append("MISSING_ACCOUNT_CURRENCY")
    if not entry_date:
        messages.append("MISSING_ENTRY_DATE")
    if account_mapping_note != "ACCOUNT_MASTER":
        messages.append(account_mapping_note)
    if amount_note == "BOTH_DEBIT_AND_CREDIT_COLUMNS_POPULATED":
        messages.append(amount_note)

    if row_looks_like_balance_or_total(record, config):
        return None, entry_date or previous_entry_date, "BALANCE_OR_TOTAL_ROW"

    record["RECORD_STATUS"] = "VALID" if not messages else "REVIEW"
    record["VALIDATION_MESSAGE"] = " | ".join(messages)
    record["RECORD_KEY"] = record_key(record)
    return record, entry_date or previous_entry_date, "OK"


def parse_balance_summary_rows(
    rows: Sequence[Sequence[Any]],
    header_idx: int,
    positions: dict[str, int],
    source_type: str,
    source_file: str,
    account_master,
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    # Real balance-summary exports essentially never have a literal "Bank
    # Group" column — resolve it the same way transaction parsing does, by
    # scanning the file's own text (name + cell content) for a known bank
    # pattern (also picks up a literal "Bank Name" cell like "CITIBANK").
    file_metadata = extract_metadata(rows[: header_idx + 2], source_file, config)
    default_bank_group = clean_text(file_metadata.get("BANK_GROUP")).upper()
    default_bank_name = clean_text(file_metadata.get("BANK_NAME"))

    for row_no, values in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(clean_text(v) for v in values):
            continue

        def get(name: str) -> Any:
            idx = positions.get(name)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        account = display_account(get("ACCOUNT_NUMBER"))
        currency = normalize_currency(get("ACCOUNT_CURRENCY"))
        statement_date = parse_date_value(get("STATEMENT_DATE"), dayfirst=True)
        opening = parse_number(get("OPENING_BALANCE"))
        closing = parse_number(get("CLOSING_BALANCE"))
        # A statement date is desirable (it drives the freshness check on the
        # dashboard) but some real exports only ever print a running/current
        # balance with no date on that specific row — don't throw the whole
        # balance away just because of that, since "balance with no date" is
        # still far more useful than no balance row at all.
        if not account or not currency or (opening is None and closing is None):
            continue

        bank_group = clean_text(get("BANK_GROUP")).upper() or default_bank_group
        bank_name = clean_text(get("BANK_NAME")) or default_bank_name
        normalized_account, entity, mapping_note = resolve_account(
            bank_group,
            account,
            currency,
            account_master,
            config,
        )
        rec = {
            "source_file_name": source_file,
            "source_type": source_type,
            "source_bank": bank_group,
            "bank_group": bank_group,
            "bank_name": bank_name,
            "account_name": clean_text(get("ACCOUNT_NAME")),
            "account_number": normalized_account,
            "currency": currency,
            "statement_date": to_excel_datetime(statement_date) if statement_date else "",
            "statement_no": clean_text(get("STATEMENT_NO")),
            "opening_balance": opening,
            "closing_balance": closing,
            "entity_code": "",
            "entity_name_short": entity,
            "entity_name": entity,
            "region": "",
            "mapping_status": "MATCHED" if mapping_note == "ACCOUNT_MASTER" else "UNMATCHED",
            "mapping_key": f"{bank_group}|{account_key(normalized_account)}|{currency}",
            "mapping_note": mapping_note,
            "page_no": None,
        }
        output.append(rec)
    return output


def balance_from_metadata_fallback(
    rows: Sequence[Sequence[Any]],
    source_type: str,
    source_file: str,
    account_master,
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    """Last-resort balance extraction for exports with no recognizable
    table at all — just scattered "Account No. / Ledger Balance" label:value
    pairs (seen on a real simple export with no transactions that period).
    Only called when neither a transaction table nor a balance-summary
    table was detected, so this never competes with either of those."""
    metadata = extract_metadata(rows, source_file, config)
    account = display_account(metadata.get("ACCOUNT_NUMBER"))
    if not account:
        return []
    opening = parse_number(metadata.get("OPENING_BALANCE"))
    closing = parse_number(metadata.get("CLOSING_BALANCE"))
    if closing is None:
        closing = parse_number(metadata.get("AVAILABLE_BALANCE"))
    if opening is None and closing is None:
        return []

    bank_group = clean_text(metadata.get("BANK_GROUP")).upper()
    bank_name = clean_text(metadata.get("BANK_NAME"))
    currency = normalize_currency(metadata.get("ACCOUNT_CURRENCY"))
    if not currency:
        currency = resolve_account_currency_hint(bank_group, account, account_master)
    normalized_account, entity, mapping_note = resolve_account(bank_group, account, currency, account_master, config)
    statement_date = parse_date_value(metadata.get("STATEMENT_DATE"), dayfirst=True)

    rec = {
        "source_file_name": source_file,
        "source_type": source_type,
        "source_bank": bank_group,
        "bank_group": bank_group,
        "bank_name": bank_name,
        "account_name": clean_text(metadata.get("ACCOUNT_NAME")),
        "account_number": normalized_account,
        "currency": currency,
        "statement_date": to_excel_datetime(statement_date) if statement_date else "",
        "statement_no": clean_text(metadata.get("STATEMENT_NO")),
        "opening_balance": opening,
        "closing_balance": closing if closing is not None else opening,
        "entity_code": "",
        "entity_name_short": entity,
        "entity_name": entity,
        "region": "",
        "mapping_status": "MATCHED" if mapping_note == "ACCOUNT_MASTER" else "UNMATCHED",
        "mapping_key": f"{bank_group}|{account_key(normalized_account)}|{currency}",
        "mapping_note": mapping_note + " | Derived from scattered label:value metadata, no transaction/balance table found",
        "page_no": None,
    }
    return [rec]


def derive_balances_from_transaction_table(
    rows: Sequence[Sequence[Any]],
    header_end: int,
    mapping: dict[str, int],
    metadata: dict[str, Any],
    source_file: str,
    source_sheet: str,
    account_master,
    config: dict[str, Any],
) -> list[dict[str, Any]]:
    """One balance row per account, derived from a TRANSACTION table.

    Many real transaction exports carry the account balance inside the same
    table instead of a separate balance-summary sheet — Citi's wide export
    repeats "Opening/Current-Closing Ledger Balance" on every row, HSBC's
    account listing repeats "Closing ledger balance", and Standard
    Chartered prints a per-transaction running "Balance" whose LAST row is
    the closing balance. The balance-summary detector must not classify
    those sheets as balance tables (they are transaction tables — see
    detect_balance_summary_header), which used to mean their balance data
    was dropped entirely: every Citi/SC account showed "no balance data" on
    the dashboard's Cash Position / Bank Debit sufficiency views.

    Per (account, currency): CLOSING_BALANCE column wins (last non-empty
    value); a running-balance column is the fallback, taking the last
    non-empty value in file order (statements are printed chronologically);
    OPENING_BALANCE takes the first non-empty value; the statement date is
    the latest entry/value date seen for that account.
    """
    pos_account = mapping.get("ACCOUNT_NUMBER")
    pos_ccy = mapping.get("ACCOUNT_CURRENCY")
    pos_open = mapping.get("OPENING_BALANCE")
    pos_close = mapping.get("CLOSING_BALANCE")
    pos_running = mapping.get("RUNNING_BALANCE")
    pos_date = mapping.get("ENTRY_DATE", mapping.get("VALUE_DATE"))
    if pos_close is None and pos_running is None:
        return []

    meta_account = display_account(metadata.get("ACCOUNT_NUMBER"))
    meta_ccy = normalize_currency(metadata.get("ACCOUNT_CURRENCY"))
    bank_group = clean_text(metadata.get("BANK_GROUP")).upper()
    bank_name = clean_text(metadata.get("BANK_NAME"))
    if not bank_group:
        # Same per-row fallback build_transaction_record() uses: the bank
        # name often only appears inside the data rows (e.g. a "CITIBANK"
        # cell on every row of the ZPI CSV), not in any title block —
        # without this, resolve_account() got an empty bank group and every
        # derived balance from such a file came out UNMAPPED even though
        # the very same file's transactions resolved fine.
        sample = source_file + "\n" + "\n".join(
            row_raw_text(r) for r in rows[header_end + 1 : header_end + 6]
        )
        inferred_group, inferred_name = infer_bank(sample, config)
        bank_group = bank_group or inferred_group
        bank_name = bank_name or inferred_name

    def cell(values: Sequence[Any], idx: int | None) -> Any:
        if idx is None or idx >= len(values):
            return None
        return values[idx]

    per_account: dict[tuple[str, str], dict[str, Any]] = {}
    for values in rows[header_end + 1 :]:
        if not any(clean_text(v) for v in values):
            continue
        account = display_account(cell(values, pos_account)) or meta_account
        currency = normalize_currency(cell(values, pos_ccy)) or meta_ccy
        if not account:
            continue
        key = (account, currency)
        agg = per_account.setdefault(key, {"opening": None, "closing": None, "running": None, "date": None, "name": ""})
        opening = parse_number(cell(values, pos_open))
        closing = parse_number(cell(values, pos_close))
        running = parse_number(cell(values, pos_running))
        if agg["opening"] is None and opening is not None:
            agg["opening"] = opening
        if closing is not None:
            agg["closing"] = closing
        if running is not None:
            agg["running"] = running
        row_date = parse_date_value(cell(values, pos_date), dayfirst=True)
        if row_date and (agg["date"] is None or row_date > agg["date"]):
            agg["date"] = row_date
        if not agg["name"]:
            agg["name"] = clean_text(cell(values, mapping.get("ACCOUNT_NAME"))) or clean_text(metadata.get("ACCOUNT_NAME"))

    output: list[dict[str, Any]] = []
    for (account, currency), agg in per_account.items():
        closing = agg["closing"] if agg["closing"] is not None else agg["running"]
        if closing is None and agg["opening"] is None:
            continue
        ccy = currency or resolve_account_currency_hint(bank_group, account, account_master)
        normalized_account, entity, mapping_note = resolve_account(bank_group, account, ccy, account_master, config)
        source = "closing-balance column" if agg["closing"] is not None else "last running-balance row"
        output.append({
            "source_file_name": source_file,
            "source_type": "TRANSACTION_TABLE_DERIVED",
            "source_bank": bank_group,
            "bank_group": bank_group,
            "bank_name": bank_name,
            "account_name": agg["name"],
            "account_number": normalized_account,
            "currency": ccy,
            "statement_date": to_excel_datetime(agg["date"]) if agg["date"] else "",
            "statement_no": "",
            "opening_balance": agg["opening"],
            "closing_balance": closing if closing is not None else agg["opening"],
            "entity_code": "",
            "entity_name_short": entity,
            "entity_name": entity,
            "region": "",
            "mapping_status": "MATCHED" if mapping_note == "ACCOUNT_MASTER" else "UNMATCHED",
            "mapping_key": f"{bank_group}|{account_key(normalized_account)}|{ccy}",
            "mapping_note": mapping_note + f" | Derived from the transaction table ({source}), sheet {source_sheet}",
            "page_no": None,
        })
    return output


def parse_transaction_matrix(
    rows: Sequence[Sequence[Any]],
    source_file: str,
    source_sheet: str,
    source_page: int | None,
    account_master,
    config: dict[str, Any],
    metadata_override: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int, list[dict[str, Any]]]:
    header = detect_transaction_header(rows, config)
    if not header:
        return [], [], 0, []
    header_start, header_end, mapping = header
    metadata = extract_metadata(rows[: max(header_end + 1, 1)], f"{source_file} {source_sheet}", config)
    # Metadata may sit below the detected header in odd exports; scan the full title area.
    metadata = {**extract_metadata(rows, f"{source_file} {source_sheet}", config), **metadata}
    if metadata_override:
        metadata = {**metadata, **{k: v for k, v in metadata_override.items() if clean_text(v)}}

    records: list[dict[str, Any]] = []
    unmapped: list[dict[str, Any]] = []
    skipped = 0
    previous_entry_date: date | None = None
    consecutive_blank = 0

    for idx in range(header_end + 1, len(rows)):
        values = list(rows[idx])
        if not any(clean_text(v) for v in values):
            consecutive_blank += 1
            if consecutive_blank >= 25:
                break
            continue
        consecutive_blank = 0
        record, previous_entry_date, reason = build_transaction_record(
            values=values,
            mapping=mapping,
            metadata=metadata,
            source_file=source_file,
            source_sheet=source_sheet,
            source_row=idx + 1,
            source_page=source_page,
            previous_entry_date=previous_entry_date,
            account_master=account_master,
            config=config,
        )
        if record:
            records.append(record)
        else:
            skipped += 1
            if reason not in {"NO_AMOUNT", "ZERO_AMOUNT", "BALANCE_OR_TOTAL_ROW"}:
                unmapped.append(
                    {
                        "SOURCE_FILE": source_file,
                        "SOURCE_SHEET": source_sheet,
                        "SOURCE_ROW": idx + 1,
                        "SOURCE_PAGE": source_page,
                        "REASON": reason,
                        "RAW_TEXT": row_raw_text(values),
                    }
                )
    # Balance data living INSIDE the transaction table (Citi/HSBC closing-
    # balance columns, SC running balance) — one derived row per account.
    derived_balances = derive_balances_from_transaction_table(
        rows, header_end, mapping, metadata, source_file, source_sheet, account_master, config
    ) if records else []
    return records, unmapped, skipped, derived_balances


def worksheet_to_rows(ws, max_rows: int | None = None, max_cols: int | None = None) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if max_rows is not None and r_idx > max_rows:
            break
        values = list(row[:max_cols] if max_cols else row)
        rows.append(values)
    return rows


def parse_excel_file(path: Path, account_master, config: dict[str, Any]) -> ParseResult:
    result = ParseResult()
    log = ParseLogEntry(source_file=path.name, source_type=path.suffix.lstrip(".").upper(), status="STARTED")
    try:
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        log.status = "ERROR"
        log.message = f"Unable to open workbook: {exc}"
        result.logs.append(log)
        return result

    try:
        for ws in wb.worksheets:
            log.sheets_or_pages_scanned += 1
            rows = worksheet_to_rows(ws)
            balance_header = detect_balance_summary_header(rows, config)
            balances = []
            if balance_header:
                header_idx, positions = balance_header
                balances = parse_balance_summary_rows(
                    rows,
                    header_idx,
                    positions,
                    source_type=path.suffix.lstrip(".").upper(),
                    source_file=path.name,
                    account_master=account_master,
                    config=config,
                )
            if balances:
                # Only treat this as "handled as a balance sheet" once we
                # actually got real balance rows out of it — a header match
                # that yields nothing (e.g. a stray data row scored as if it
                # were a header) must still fall through to transaction
                # parsing below instead of silently discarding the sheet.
                result.balances.extend(balances)
                log.balance_rows += len(balances)
                continue

            records, unmapped, skipped, derived_balances = parse_transaction_matrix(
                rows,
                source_file=path.name,
                source_sheet=ws.title,
                source_page=None,
                account_master=account_master,
                config=config,
            )
            if not records:
                # Neither a transaction table nor a balance-summary table was
                # found on this sheet — try scattered label:value metadata
                # (e.g. a plain "Account No. / Ledger Balance" snippet with
                # no transactions that period) before giving up on it.
                fallback_balances = balance_from_metadata_fallback(
                    rows,
                    source_type=path.suffix.lstrip(".").upper(),
                    source_file=path.name,
                    account_master=account_master,
                    config=config,
                )
                if fallback_balances:
                    result.balances.extend(fallback_balances)
                    log.balance_rows += len(fallback_balances)
                    continue
            result.transactions.extend(records)
            result.unmapped_rows.extend(unmapped)
            log.transaction_rows += len(records)
            log.review_rows += sum(1 for rec in records if rec.get("RECORD_STATUS") != "VALID")
            log.skipped_rows += skipped
            if derived_balances:
                result.balances.extend(derived_balances)
                log.balance_rows += len(derived_balances)
    finally:
        wb.close()

    bank_group, bank_name = infer_bank(path.name, config)
    log.bank_group = bank_group
    log.bank_name = bank_name
    if log.transaction_rows or log.balance_rows:
        log.status = "PARSED"
        log.message = "Transaction detail and/or balance summary detected."
    else:
        log.status = "NO_DATA"
        log.message = "No recognized transaction table or balance summary found."
    result.logs.append(log)
    return result


def read_delimited_rows(path: Path) -> list[list[Any]]:
    raw = path.read_bytes()

    # cp1252/latin-1 never raise on any byte sequence, so a UTF-16 file
    # (every ASCII char followed by a 0x00 byte) would otherwise "succeed"
    # under one of those, decoding every character as itself plus a stray
    # NUL and silently mangling the whole file into one-space-per-letter
    # text (seen on a real export with no BOM at all). Detect that pattern
    # up front — via BOM if present, else by counting NUL bytes in each
    # half of a sample — and put the right UTF-16 variant first.
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]
    sample_bytes = raw[:400]
    if sample_bytes[:2] == b"\xff\xfe":
        encodings = ["utf-16-le"] + encodings
    elif sample_bytes[:2] == b"\xfe\xff":
        encodings = ["utf-16-be"] + encodings
    elif len(sample_bytes) >= 20:
        even_nulls = sample_bytes[0::2].count(0)
        odd_nulls = sample_bytes[1::2].count(0)
        if max(even_nulls, odd_nulls) > len(sample_bytes) * 0.3:
            encodings = (["utf-16-le"] if odd_nulls >= even_nulls else ["utf-16-be"]) + encodings

    last_error: Exception | None = None
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            sample = text[:10000]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                delimiter = dialect.delimiter
            except csv.Error:
                delimiter = "," if path.suffix.lower() == ".csv" else "\t"
            return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]
        except Exception as exc:
            last_error = exc
    raise ParserError(f"Unable to read delimited file {path}: {last_error}")


def parse_delimited_file(path: Path, account_master, config: dict[str, Any]) -> ParseResult:
    result = ParseResult()
    log = ParseLogEntry(source_file=path.name, source_type=path.suffix.lstrip(".").upper(), status="STARTED")
    try:
        rows = read_delimited_rows(path)
        log.sheets_or_pages_scanned = 1
        balance_header = detect_balance_summary_header(rows, config)
        balances = []
        if balance_header:
            header_idx, positions = balance_header
            balances = parse_balance_summary_rows(
                rows,
                header_idx,
                positions,
                source_type=path.suffix.lstrip(".").upper(),
                source_file=path.name,
                account_master=account_master,
                config=config,
            )
        if balances:
            result.balances.extend(balances)
            log.balance_rows = len(balances)
        else:
            records, unmapped, skipped, derived_balances = parse_transaction_matrix(
                rows,
                source_file=path.name,
                source_sheet="Delimited",
                source_page=None,
                account_master=account_master,
                config=config,
            )
            if not records:
                fallback_balances = balance_from_metadata_fallback(
                    rows,
                    source_type=path.suffix.lstrip(".").upper(),
                    source_file=path.name,
                    account_master=account_master,
                    config=config,
                )
                if fallback_balances:
                    result.balances.extend(fallback_balances)
                    log.balance_rows = len(fallback_balances)
            if not log.balance_rows:
                result.transactions.extend(records)
                result.unmapped_rows.extend(unmapped)
                log.transaction_rows = len(records)
                log.review_rows = sum(1 for rec in records if rec.get("RECORD_STATUS") != "VALID")
                log.skipped_rows = skipped
                if derived_balances:
                    result.balances.extend(derived_balances)
                    log.balance_rows = len(derived_balances)
        log.status = "PARSED" if log.transaction_rows or log.balance_rows else "NO_DATA"
        log.message = "Delimited source parsed." if log.status == "PARSED" else "No recognized table found."
    except Exception as exc:
        log.status = "ERROR"
        log.message = str(exc)
    result.logs.append(log)
    return result


# Deutsche Bank's "db-direct internet / Balance And Transaction Detail
# Report" export is free-flowing narrative text (label: value pairs spread
# across several lines per transaction), not a grid pdfplumber's
# extract_tables() can find — hence every page previously landed in the
# unmapped/review queue. One page = one account + one statement date.
DB_REPORT_SIGNATURE_RE = re.compile(r"db-direct internet|Balance And Transaction Detail Report", re.IGNORECASE)
DB_ACCOUNT_LINE_RE = re.compile(r"Account Number:\s*(\S+)\s+Statement Date:\s*(\d{2}\.\d{2}\.\d{4})")
DB_ACCOUNT_NAME_RE = re.compile(r"Account Name:\s*(.+)")
DB_BANK_BRANCH_RE = re.compile(r"Bank Branch:\s*(.+)")
DB_CURRENCY_RE = re.compile(r"^Currency:\s*([A-Z]{3})")
DB_CLOSING_BALANCE_RE = re.compile(r"Closing Book Balance:\s*([\d,()\.\-]+)")
# A transaction block starts with "<amount> <book date> Transaction Type: <type>".
DB_TXN_START_RE = re.compile(r"^(-?[\d,]+\.\d{2})\s+(\d{2}\.\d{2}\.\d{4})\s+Transaction Type:\s*(.*)$")
DB_VALUE_DATE_RE = re.compile(r"^(\d{2}\.\d{2}\.\d{4})\s+Customer Reference:")
DB_CUSTOMER_REF_RE = re.compile(r"Customer Reference:\s*(\S+)")
DB_BANK_REF_RE = re.compile(r"Bank Reference:\s*(\S+)")
DB_DETAILS_RE = re.compile(
    r"Details of Payment:\s*(.*?)(?=(?:\n-?[\d,]+\.\d{2}\s+\d{2}\.\d{2}\.\d{4}\s+Transaction Type)|(?:\n-?[\d,]+\.\d{2}\s+TOTAL)|\Z)",
    re.DOTALL,
)


def parse_deutsche_bank_pdf_page(
    text: str,
    source_file: str,
    page_no: int,
    account_master,
    config: dict[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    """Returns (transaction_records, balance_record_or_None) for one page of
    a db-direct internet report, or ([], None) if this page doesn't match
    that report's signature at all (so the caller can fall back to generic
    table parsing)."""
    if not DB_REPORT_SIGNATURE_RE.search(text):
        return [], None

    account_match = DB_ACCOUNT_LINE_RE.search(text)
    if not account_match:
        return [], None
    account_number, statement_date_raw = account_match.groups()
    statement_date = parse_date_value(statement_date_raw, dayfirst=True)

    name_match = DB_ACCOUNT_NAME_RE.search(text)
    account_name = clean_text(name_match.group(1)) if name_match else ""
    branch_match = DB_BANK_BRANCH_RE.search(text)
    # Deutsche Bank statements are branch-specific (Singapore, Ho Chi Minh,
    # ...) — read the real branch per page instead of assuming Singapore.
    bank_name = clean_text(branch_match.group(1)) if branch_match else "DEUTSCHE BANK"
    currency = ""
    for line in text.splitlines():
        m = DB_CURRENCY_RE.match(line.strip())
        if m:
            currency = m.group(1)
            break
    closing_match = DB_CLOSING_BALANCE_RE.search(text)
    # SUMMARY lines can show one amount per currency (e.g. "SGD USD @ rate")
    # — the first number is always the account's own (native) currency.
    closing_balance = parse_number(closing_match.group(1).split()[0]) if closing_match else None

    normalized_account, entity, mapping_note = resolve_account("DB", account_number, currency, account_master, config)

    lines = text.splitlines()
    starts = [i for i, line in enumerate(lines) if DB_TXN_START_RE.match(line.strip())]
    records: list[dict[str, Any]] = []
    for k, start_idx in enumerate(starts):
        end_idx = starts[k + 1] if k + 1 < len(starts) else len(lines)
        block_lines = lines[start_idx:end_idx]
        block_text = "\n".join(block_lines)
        m = DB_TXN_START_RE.match(block_lines[0].strip())
        amount = parse_number(m.group(1))
        book_date_raw = m.group(2)
        txn_type = clean_text(m.group(3))

        value_date_match = DB_VALUE_DATE_RE.search(block_text)
        value_date_raw = value_date_match.group(1) if value_date_match else book_date_raw
        customer_ref_match = DB_CUSTOMER_REF_RE.search(block_text)
        customer_ref = clean_text(customer_ref_match.group(1)) if customer_ref_match else ""
        bank_ref_match = DB_BANK_REF_RE.search(block_text)
        bank_ref = clean_text(bank_ref_match.group(1)) if bank_ref_match else ""
        details_match = DB_DETAILS_RE.search(block_text)
        details = clean_text(details_match.group(1)) if details_match else ""

        entry_date = parse_date_value(book_date_raw, dayfirst=True)
        value_date = parse_date_value(value_date_raw, dayfirst=True)
        eform = detect_eform([customer_ref, bank_ref, details], config)

        rec: dict[str, Any] = {col: "" for col in TRANSACTION_OUTPUT_COLUMNS}
        rec.update(
            {
                "BANK_NAME": bank_name,
                "ACCOUNT_NUMBER": normalized_account,
                "ACCOUNT_CURRENCY": currency,
                "BANK_REFERENCE": bank_ref,
                "CUSTOMER_REFERENCE": customer_ref,
                "ENTRY_DATE": to_excel_datetime(entry_date) if entry_date else "",
                "TRANSACTION_AMOUNT": amount,
                "DEBIT_CREDIT_INDICATOR": "DEBIT" if (amount or 0) < 0 else "CREDIT",
                "PRODUCT_TYPE": txn_type,
                "TRANSACTION_DESCRIPTION": details,
                "PAYMENT_DETAILS": details,
                "SOURCE_FILE": source_file,
                "SOURCE_SHEET": f"Page {page_no}",
                "BANK_GROUP": "DB",
                "ACCOUNT_NAME": account_name,
                "STATEMENT_DATE": to_excel_datetime(statement_date) if statement_date else "",
                "STATEMENT_NO": "",
                "VALUE_DATE": to_excel_datetime(value_date) if value_date else "",
                "ENTITY": entity,
                "SOURCE_ROW": start_idx + 1,
                "SOURCE_PAGE": page_no,
                "RAW_TEXT": clean_text(block_text)[:4000],
                "EFORM_CANDIDATE": eform,
            }
        )
        messages = []
        if mapping_note != "ACCOUNT_MASTER":
            messages.append(mapping_note)
        rec["RECORD_STATUS"] = "VALID" if not messages else "REVIEW"
        rec["VALIDATION_MESSAGE"] = " | ".join(messages)
        rec["RECORD_KEY"] = record_key(rec)
        records.append(rec)

    balance_rec = None
    if closing_balance is not None:
        balance_rec = {
            "source_file_name": source_file,
            "source_type": "PDF",
            "source_bank": "DB",
            "bank_group": "DB",
            "bank_name": bank_name,
            "account_name": account_name,
            "account_number": normalized_account,
            "currency": currency,
            "statement_date": to_excel_datetime(statement_date) if statement_date else "",
            "statement_no": "",
            "opening_balance": None,
            "closing_balance": closing_balance,
            "entity_code": "",
            "entity_name_short": entity,
            "entity_name": entity,
            "region": "",
            "mapping_status": "MATCHED" if mapping_note == "ACCOUNT_MASTER" else "UNMATCHED",
            "mapping_key": f"DB|{account_key(normalized_account)}|{currency}",
            "mapping_note": mapping_note,
            "page_no": page_no,
        }
    return records, balance_rec


def normalize_pdf_table(table: Sequence[Sequence[Any]]) -> list[list[Any]]:
    width = max((len(row) for row in table if row), default=0)
    rows: list[list[Any]] = []
    for row in table:
        vals = list(row or [])
        if len(vals) < width:
            vals.extend([None] * (width - len(vals)))
        rows.append(vals)
    return rows


def parse_pdf_file(path: Path, account_master, config: dict[str, Any]) -> ParseResult:
    result = ParseResult()
    log = ParseLogEntry(source_file=path.name, source_type="PDF", status="STARTED")
    if pdfplumber is None:
        log.status = "ERROR"
        log.message = "pdfplumber is not installed. Run: pip install pdfplumber"
        result.logs.append(log)
        return result

    try:
        with pdfplumber.open(path) as pdf:
            combined_metadata_text = path.name
            carried_metadata: dict[str, Any] = {}
            for page_number, page in enumerate(pdf.pages, start=1):
                log.sheets_or_pages_scanned += 1
                page_text = page.extract_text() or ""
                combined_metadata_text += "\n" + page_text
                page_metadata_rows = [[line] for line in page_text.splitlines() if clean_text(line)]
                page_metadata = extract_metadata(page_metadata_rows, path.name, config) if page_metadata_rows else {}
                carried_metadata = {**carried_metadata, **{k: v for k, v in page_metadata.items() if clean_text(v)}}
                page_found = False

                db_records, db_balance = parse_deutsche_bank_pdf_page(
                    page_text, path.name, page_number, account_master, config
                )
                if db_records or db_balance:
                    page_found = True
                    result.transactions.extend(db_records)
                    log.transaction_rows += len(db_records)
                    log.review_rows += sum(1 for rec in db_records if rec.get("RECORD_STATUS") != "VALID")
                    if db_balance:
                        result.balances.append(db_balance)
                        log.balance_rows += 1
                    continue

                tables = page.extract_tables() or []
                for table_number, table in enumerate(tables, start=1):
                    rows = normalize_pdf_table(table)
                    if not rows:
                        continue
                    records, unmapped, skipped, derived_balances = parse_transaction_matrix(
                        rows,
                        source_file=path.name,
                        source_sheet=f"Page {page_number} Table {table_number}",
                        source_page=page_number,
                        account_master=account_master,
                        config=config,
                        metadata_override=carried_metadata,
                    )
                    if records:
                        page_found = True
                    result.transactions.extend(records)
                    result.unmapped_rows.extend(unmapped)
                    log.transaction_rows += len(records)
                    log.review_rows += sum(1 for rec in records if rec.get("RECORD_STATUS") != "VALID")
                    log.skipped_rows += skipped
                    if derived_balances:
                        result.balances.extend(derived_balances)
                        log.balance_rows += len(derived_balances)
                if not page_found and page_text:
                    # No generic table could be recognized. Keep the page in the review queue.
                    result.unmapped_rows.append(
                        {
                            "SOURCE_FILE": path.name,
                            "SOURCE_SHEET": f"Page {page_number}",
                            "SOURCE_ROW": "",
                            "SOURCE_PAGE": page_number,
                            "REASON": "PDF_PAGE_WITHOUT_RECOGNIZED_TRANSACTION_TABLE",
                            "RAW_TEXT": clean_text(page_text)[:32000],
                        }
                    )
            bank_group, bank_name = infer_bank(combined_metadata_text, config)
            log.bank_group = bank_group
            log.bank_name = bank_name
    except Exception as exc:
        log.status = "ERROR"
        log.message = f"Unable to parse PDF: {exc}"
        result.logs.append(log)
        return result

    if log.transaction_rows or log.balance_rows:
        log.status = "PARSED"
        log.message = "PDF transaction table and/or balance detected."
    else:
        log.status = "REVIEW"
        log.message = (
            "No transaction table was recognized in the PDF. The PDF may be scanned or require a bank-specific adapter."
        )
    result.logs.append(log)
    return result


def list_input_files(input_path: Path, recursive: bool, excluded: set[Path]) -> list[Path]:
    if input_path.is_file():
        candidates = [input_path]
    elif input_path.is_dir():
        iterator = input_path.rglob("*") if recursive else input_path.glob("*")
        candidates = [p for p in iterator if p.is_file()]
    else:
        raise FileNotFoundError(f"Input path not found: {input_path}")

    files: list[Path] = []
    generated_names = {
        "bank_transaction_detail.xlsx",
        "bank_transaction_detail.csv",
        "summary_rebuilt.xlsx",
    }
    for path in candidates:
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        if path.name.lower() in generated_names:
            continue
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        resolved = path.resolve()
        if any(resolved == ex.resolve() for ex in excluded if ex.exists()):
            continue
        files.append(path)
    return sorted(files, key=lambda p: str(p).lower())


def merge_results(target: ParseResult, source: ParseResult) -> None:
    target.transactions.extend(source.transactions)
    target.balances.extend(source.balances)
    target.logs.extend(source.logs)
    target.unmapped_rows.extend(source.unmapped_rows)
    target.duplicates.extend(source.duplicates)


def deduplicate_transactions(result: ParseResult) -> None:
    unique: dict[str, dict[str, Any]] = {}
    for rec in result.transactions:
        key = clean_text(rec.get("RECORD_KEY")) or record_key(rec)
        rec["RECORD_KEY"] = key
        if key in unique:
            duplicate = dict(rec)
            duplicate["DUPLICATE_OF_SOURCE_FILE"] = unique[key].get("SOURCE_FILE")
            duplicate["DUPLICATE_OF_SOURCE_ROW"] = unique[key].get("SOURCE_ROW")
            result.duplicates.append(duplicate)
        else:
            unique[key] = rec
    result.transactions = list(unique.values())

    duplicate_counts = defaultdict(int)
    for rec in result.duplicates:
        duplicate_counts[clean_text(rec.get("SOURCE_FILE"))] += 1
    for log in result.logs:
        log.duplicate_rows = duplicate_counts[log.source_file]


def balance_record_key(rec: dict[str, Any]) -> str:
    stmt_date = parse_date_value(rec.get("statement_date"))
    return "|".join(
        [
            clean_text(rec.get("bank_group")).upper(),
            account_key(rec.get("account_number")),
            normalize_currency(rec.get("currency")),
            stmt_date.isoformat() if stmt_date else "",
            clean_text(rec.get("statement_no")),
            f"{parse_number(rec.get('opening_balance')) or 0:.6f}",
            f"{parse_number(rec.get('closing_balance')) or 0:.6f}",
        ]
    )


def deduplicate_balances(result: ParseResult) -> None:
    unique: dict[str, dict[str, Any]] = {}
    for rec in result.balances:
        unique[balance_record_key(rec)] = rec
    result.balances = list(unique.values())


def parse_all(
    input_path: Path,
    recursive: bool,
    config: dict[str, Any],
    config_path: Path | None,
    output_path: Path,
    summary_output_path: Path,
) -> ParseResult:
    base_dir = config_path.parent if config_path else Path.cwd()
    master_setting = clean_text(config.get("account_master_file"))
    account_master_path = Path(master_setting) if master_setting else None
    if account_master_path and not account_master_path.is_absolute():
        account_master_path = base_dir / account_master_path
    account_master = load_account_master(account_master_path)

    files = list_input_files(
        input_path,
        recursive=recursive,
        excluded={output_path, summary_output_path},
    )
    if not files:
        raise ParserError(f"No supported files found under: {input_path}")

    aggregate = ParseResult()
    for path in files:
        logging.info("Parsing %s", path)
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            result = parse_excel_file(path, account_master, config)
        elif suffix in {".csv", ".txt"}:
            result = parse_delimited_file(path, account_master, config)
        elif suffix == ".pdf":
            result = parse_pdf_file(path, account_master, config)
        else:
            continue
        merge_results(aggregate, result)

    deduplicate_transactions(aggregate)
    deduplicate_balances(aggregate)
    aggregate.transactions.sort(
        key=lambda rec: (
            parse_date_value(rec.get("ENTRY_DATE")) or date.min,
            clean_text(rec.get("BANK_GROUP")),
            account_key(rec.get("ACCOUNT_NUMBER")),
            parse_number(rec.get("TRANSACTION_AMOUNT")) or 0,
        )
    )
    aggregate.balances.sort(
        key=lambda rec: (
            parse_date_value(rec.get("statement_date")) or date.min,
            clean_text(rec.get("bank_group")),
            account_key(rec.get("account_number")),
        )
    )
    return aggregate


def sheet_title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1).value = title
    ws.cell(1, 1).font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1).value = subtitle
    ws.cell(2, 1).font = Font(name="Aptos", size=9, color=WHITE)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def style_header(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def add_table(ws, name: str, header_row: int, columns: int, last_row: int) -> None:
    if last_row <= header_row:
        return
    ref = f"A{header_row}:{get_column_letter(columns)}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def set_column_widths(ws, widths: dict[str, float], default: float = 13) -> None:
    for idx in range(1, ws.max_column + 1):
        header = clean_text(ws.cell(5, idx).value) if ws.max_row >= 5 else clean_text(ws.cell(1, idx).value)
        width = widths.get(header, default)
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_transaction_sheet(wb: Workbook, records: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("TRANSACTION_DETAIL")
    sheet_title(
        ws,
        "Bank Transaction Detail — Parser Output",
        "One row = one bank transaction. Debit amount is negative; credit amount is positive. The first 16 columns are compatible with Treasury Payment Control Tower v6.",
        len(TRANSACTION_OUTPUT_COLUMNS),
    )
    header_row = 5
    for col, header in enumerate(TRANSACTION_OUTPUT_COLUMNS, start=1):
        ws.cell(header_row, col).value = header
    style_header(ws, header_row, len(TRANSACTION_OUTPUT_COLUMNS))

    for row_idx, record in enumerate(records, start=header_row + 1):
        for col_idx, header in enumerate(TRANSACTION_OUTPUT_COLUMNS, start=1):
            cell = ws.cell(row_idx, col_idx)
            value = xml_safe(record.get(header))
            cell.value = value
            cell.font = Font(name="Aptos", size=9, color=GREEN)
            cell.alignment = Alignment(vertical="top", wrap_text=header in {"TRANSACTION_DESCRIPTION", "EXTRA_INFORMATION", "PAYMENT_DETAILS", "RAW_TEXT", "VALIDATION_MESSAGE"})
            if header in {"ENTRY_DATE", "STATEMENT_DATE", "VALUE_DATE"} and value:
                cell.number_format = DATE_FORMAT
            elif header == "TRANSACTION_AMOUNT":
                cell.number_format = AMOUNT_FORMAT
            elif header in {"ACCOUNT_NUMBER", "BENEFICIARY_ACCOUNT", "BANK_REFERENCE", "CUSTOMER_REFERENCE", "BANK_TRANSACTION_ID"}:
                cell.number_format = "@"

    last_row = max(header_row + 1, header_row + len(records))
    add_table(ws, "tblTransactionDetail", header_row, len(TRANSACTION_OUTPUT_COLUMNS), last_row)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(TRANSACTION_OUTPUT_COLUMNS))}{last_row}"
    ws.sheet_view.showGridLines = False

    widths = {
        "BANK_NAME": 24,
        "ACCOUNT_NUMBER": 18,
        "ACCOUNT_CURRENCY": 12,
        "BANK_REFERENCE": 24,
        "CUSTOMER_REFERENCE": 24,
        "ENTRY_DATE": 12,
        "TRANSACTION_AMOUNT": 18,
        "DEBIT_CREDIT_INDICATOR": 14,
        "PRODUCT_TYPE": 22,
        "TRANSACTION_DESCRIPTION": 40,
        "EXTRA_INFORMATION": 36,
        "PAYMENT_DETAILS": 40,
        "BENEFICIARY_REFERENCE": 24,
        "REMITTER_REFERENCE": 24,
        "SOURCE_FILE": 30,
        "SOURCE_SHEET": 22,
        "BANK_GROUP": 12,
        "ACCOUNT_NAME": 28,
        "STATEMENT_DATE": 14,
        "STATEMENT_NO": 16,
        "VALUE_DATE": 12,
        "BENEFICIARY_NAME": 30,
        "BENEFICIARY_ACCOUNT": 22,
        "END_TO_END_REFERENCE": 26,
        "BANK_TRANSACTION_ID": 26,
        "EFORM_CANDIDATE": 28,
        "ENTITY": 18,
        "SOURCE_ROW": 12,
        "SOURCE_PAGE": 12,
        "RAW_TEXT": 50,
        "RECORD_STATUS": 14,
        "VALIDATION_MESSAGE": 36,
        "RECORD_KEY": 26,
    }
    set_column_widths(ws, widths)

    if records:
        status_col = TRANSACTION_OUTPUT_COLUMNS.index("RECORD_STATUS") + 1
        status_letter = get_column_letter(status_col)
        data_range = f"A{header_row + 1}:{get_column_letter(len(TRANSACTION_OUTPUT_COLUMNS))}{last_row}"
        ws.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[f'${status_letter}{header_row + 1}="REVIEW"'],
                fill=PatternFill("solid", fgColor=LIGHT_ORANGE),
            ),
        )


def write_balance_sheet(wb: Workbook, balances: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("BALANCE_SUMMARY")
    sheet_title(
        ws,
        "Bank Balance Summary",
        "Balance-level records retained for liquidity and statement coverage. These rows are not used to confirm an individual payment debit.",
        len(BALANCE_OUTPUT_COLUMNS),
    )
    header_row = 5
    for col, header in enumerate(BALANCE_OUTPUT_COLUMNS, start=1):
        ws.cell(header_row, col).value = header
    style_header(ws, header_row, len(BALANCE_OUTPUT_COLUMNS))

    for row_idx, record in enumerate(balances, start=header_row + 1):
        for col_idx, header in enumerate(BALANCE_OUTPUT_COLUMNS, start=1):
            cell = ws.cell(row_idx, col_idx)
            cell.value = xml_safe(record.get(header))
            cell.font = Font(name="Aptos", size=9, color=GREEN)
            if header == "statement_date" and cell.value:
                cell.number_format = DATE_FORMAT
            elif header in {"opening_balance", "closing_balance"}:
                cell.number_format = AMOUNT_FORMAT
            elif header == "account_number":
                cell.number_format = "@"

    last_row = max(header_row + 1, header_row + len(balances))
    add_table(ws, "tblBalanceSummary", header_row, len(BALANCE_OUTPUT_COLUMNS), last_row)
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    widths = {header: 16 for header in BALANCE_OUTPUT_COLUMNS}
    widths.update(
        {
            "source_file_name": 30,
            "bank_name": 30,
            "account_name": 30,
            "account_number": 18,
            "statement_date": 14,
            "opening_balance": 18,
            "closing_balance": 18,
            "entity_name": 28,
            "mapping_note": 32,
        }
    )
    set_column_widths(ws, widths)


def build_balance_recon(
    transactions: list[dict[str, Any]], balances: list[dict[str, Any]], tolerance: float = 0.01
) -> list[dict[str, Any]]:
    tx_by_key: dict[tuple[str, str, str, date], list[dict[str, Any]]] = defaultdict(list)
    for tx in transactions:
        entry = parse_date_value(tx.get("ENTRY_DATE"))
        if not entry:
            continue
        key = (
            clean_text(tx.get("BANK_GROUP")).upper(),
            account_key(tx.get("ACCOUNT_NUMBER")),
            normalize_currency(tx.get("ACCOUNT_CURRENCY")),
            entry,
        )
        tx_by_key[key].append(tx)

    rows: list[dict[str, Any]] = []
    for bal in balances:
        stmt_date = parse_date_value(bal.get("statement_date"))
        if not stmt_date:
            continue
        key = (
            clean_text(bal.get("bank_group")).upper(),
            account_key(bal.get("account_number")),
            normalize_currency(bal.get("currency")),
            stmt_date,
        )
        txs = tx_by_key.get(key, [])
        total_debit = sum(abs(parse_number(tx.get("TRANSACTION_AMOUNT")) or 0) for tx in txs if (parse_number(tx.get("TRANSACTION_AMOUNT")) or 0) < 0)
        total_credit = sum(parse_number(tx.get("TRANSACTION_AMOUNT")) or 0 for tx in txs if (parse_number(tx.get("TRANSACTION_AMOUNT")) or 0) > 0)
        opening = parse_number(bal.get("opening_balance")) or 0.0
        closing = parse_number(bal.get("closing_balance")) or 0.0
        calculated = opening + total_credit - total_debit
        diff = closing - calculated
        status = "PASS" if abs(diff) <= tolerance else ("NO_TRANSACTION_DETAIL" if not txs else "DIFF")
        rows.append(
            {
                "BANK_GROUP": key[0],
                "ACCOUNT_NUMBER": bal.get("account_number"),
                "CURRENCY": key[2],
                "STATEMENT_DATE": to_excel_datetime(stmt_date),
                "OPENING_BALANCE": opening,
                "TOTAL_CREDIT": total_credit,
                "TOTAL_DEBIT": total_debit,
                "CALCULATED_CLOSING": calculated,
                "REPORTED_CLOSING": closing,
                "BALANCE_DIFF": diff,
                "TRANSACTION_COUNT": len(txs),
                "RECON_STATUS": status,
                "SOURCE_FILE": bal.get("source_file_name"),
            }
        )
    return rows


def write_balance_recon_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    headers = [
        "BANK_GROUP",
        "ACCOUNT_NUMBER",
        "CURRENCY",
        "STATEMENT_DATE",
        "OPENING_BALANCE",
        "TOTAL_CREDIT",
        "TOTAL_DEBIT",
        "CALCULATED_CLOSING",
        "REPORTED_CLOSING",
        "BALANCE_DIFF",
        "TRANSACTION_COUNT",
        "RECON_STATUS",
        "SOURCE_FILE",
    ]
    ws = wb.create_sheet("BALANCE_RECON")
    sheet_title(
        ws,
        "Balance Reconciliation",
        "Check: Opening Balance + Total Credit - Total Debit = Reported Closing Balance. NO_TRANSACTION_DETAIL means the source contained only balances or transaction rows were not recognized.",
        len(headers),
    )
    header_row = 5
    for col, header in enumerate(headers, 1):
        ws.cell(header_row, col).value = header
    style_header(ws, header_row, len(headers))
    for row_idx, record in enumerate(rows, header_row + 1):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row_idx, col_idx)
            cell.value = xml_safe(record.get(header))
            cell.font = Font(name="Aptos", size=9, color=BLACK)
            if header == "STATEMENT_DATE" and cell.value:
                cell.number_format = DATE_FORMAT
            elif header in {"OPENING_BALANCE", "TOTAL_CREDIT", "TOTAL_DEBIT", "CALCULATED_CLOSING", "REPORTED_CLOSING", "BALANCE_DIFF"}:
                cell.number_format = AMOUNT_FORMAT
    last_row = max(header_row + 1, header_row + len(rows))
    add_table(ws, "tblBalanceRecon", header_row, len(headers), last_row)
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    widths = {h: 18 for h in headers}
    widths["ACCOUNT_NUMBER"] = 20
    widths["SOURCE_FILE"] = 30
    widths["RECON_STATUS"] = 22
    set_column_widths(ws, widths)


def write_quality_sheet(wb: Workbook, result: ParseResult) -> None:
    ws = wb.create_sheet("QUALITY_CHECK")
    sheet_title(
        ws,
        "Parser Quality Check",
        "Review this sheet before importing transaction detail into the Treasury dashboard.",
        8,
    )
    metrics = [
        ("Parser version", APP_VERSION),
        ("Generated at", datetime.now()),
        ("Files scanned", len(result.logs)),
        ("Transaction rows retained", len(result.transactions)),
        ("Valid transaction rows", sum(1 for rec in result.transactions if rec.get("RECORD_STATUS") == "VALID")),
        ("Transaction rows requiring review", sum(1 for rec in result.transactions if rec.get("RECORD_STATUS") != "VALID")),
        ("Exact duplicates removed", len(result.duplicates)),
        ("Balance rows retained", len(result.balances)),
        ("Unmapped/raw review rows", len(result.unmapped_rows)),
        ("Debit total", sum(abs(parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) for rec in result.transactions if (parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) < 0)),
        ("Credit total", sum((parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) for rec in result.transactions if (parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) > 0)),
    ]
    ws["A4"] = "CONTROL METRIC"
    ws["B4"] = "VALUE"
    style_header(ws, 4, 2)
    for idx, (label, value) in enumerate(metrics, start=5):
        ws.cell(idx, 1).value = label
        ws.cell(idx, 1).font = Font(name="Aptos", size=10, color=GRAY)
        ws.cell(idx, 2).value = value
        ws.cell(idx, 2).font = Font(name="Aptos", size=10, color=BLACK)
        if isinstance(value, datetime):
            ws.cell(idx, 2).number_format = DATETIME_FORMAT
        if label in {"Debit total", "Credit total"}:
            ws.cell(idx, 2).number_format = AMOUNT_FORMAT

    start_row = 18
    headers = ["BANK_GROUP", "ACCOUNT_NUMBER", "CURRENCY", "TRANSACTION_COUNT", "DEBIT_TOTAL", "CREDIT_TOTAL", "REVIEW_COUNT", "LATEST_ENTRY_DATE"]
    for col, header in enumerate(headers, 1):
        ws.cell(start_row, col).value = header
    style_header(ws, start_row, len(headers))

    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for rec in result.transactions:
        grouped[
            (
                clean_text(rec.get("BANK_GROUP")).upper(),
                display_account(rec.get("ACCOUNT_NUMBER")),
                normalize_currency(rec.get("ACCOUNT_CURRENCY")),
            )
        ].append(rec)

    for row_idx, (key, records) in enumerate(sorted(grouped.items()), start=start_row + 1):
        dates = [parse_date_value(rec.get("ENTRY_DATE")) for rec in records]
        values = [
            key[0],
            key[1],
            key[2],
            len(records),
            sum(abs(parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) for rec in records if (parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) < 0),
            sum(parse_number(rec.get("TRANSACTION_AMOUNT")) or 0 for rec in records if (parse_number(rec.get("TRANSACTION_AMOUNT")) or 0) > 0),
            sum(1 for rec in records if rec.get("RECORD_STATUS") != "VALID"),
            to_excel_datetime(max((d for d in dates if d), default=None)),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col).value = xml_safe(value)
            ws.cell(row_idx, col).font = Font(name="Aptos", size=9, color=BLACK)
        ws.cell(row_idx, 5).number_format = AMOUNT_FORMAT
        ws.cell(row_idx, 6).number_format = AMOUNT_FORMAT
        if values[7]:
            ws.cell(row_idx, 8).number_format = DATE_FORMAT

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    for col in range(3, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.sheet_view.showGridLines = False


def write_log_sheet(wb: Workbook, logs: list[ParseLogEntry]) -> None:
    headers = [
        "SOURCE_FILE",
        "SOURCE_TYPE",
        "STATUS",
        "BANK_GROUP",
        "BANK_NAME",
        "SHEETS_OR_PAGES_SCANNED",
        "TRANSACTION_ROWS",
        "BALANCE_ROWS",
        "SKIPPED_ROWS",
        "REVIEW_ROWS",
        "DUPLICATE_ROWS",
        "MESSAGE",
    ]
    ws = wb.create_sheet("PARSER_LOG")
    sheet_title(
        ws,
        "Parser Execution Log",
        "Every input file is listed here. REVIEW or NO_DATA normally means a bank-specific mapping or a raw statement sample is required.",
        len(headers),
    )
    header_row = 5
    for col, header in enumerate(headers, 1):
        ws.cell(header_row, col).value = header
    style_header(ws, header_row, len(headers))
    for row_idx, log in enumerate(logs, start=header_row + 1):
        values = [
            log.source_file,
            log.source_type,
            log.status,
            log.bank_group,
            log.bank_name,
            log.sheets_or_pages_scanned,
            log.transaction_rows,
            log.balance_rows,
            log.skipped_rows,
            log.review_rows,
            log.duplicate_rows,
            log.message,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col).value = xml_safe(value)
            ws.cell(row_idx, col).font = Font(name="Aptos", size=9, color=BLACK)
            ws.cell(row_idx, col).alignment = Alignment(vertical="top", wrap_text=col == 12)
    last_row = max(header_row + 1, header_row + len(logs))
    add_table(ws, "tblParserLog", header_row, len(headers), last_row)
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    widths = {h: 18 for h in headers}
    widths.update({"SOURCE_FILE": 32, "BANK_NAME": 32, "STATUS": 14, "MESSAGE": 58})
    set_column_widths(ws, widths)


def write_simple_records_sheet(
    wb: Workbook,
    sheet_name: str,
    title: str,
    subtitle: str,
    records: list[dict[str, Any]],
    preferred_headers: list[str],
    table_name: str,
) -> None:
    headers = list(preferred_headers)
    extra = sorted({key for rec in records for key in rec.keys()} - set(headers))
    headers.extend(extra)
    if not headers:
        headers = ["MESSAGE"]
    ws = wb.create_sheet(sheet_name)
    sheet_title(ws, title, subtitle, len(headers))
    header_row = 5
    for col, header in enumerate(headers, 1):
        ws.cell(header_row, col).value = header
    style_header(ws, header_row, len(headers))
    for row_idx, rec in enumerate(records, start=header_row + 1):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row_idx, col_idx).value = xml_safe(rec.get(header))
            ws.cell(row_idx, col_idx).font = Font(name="Aptos", size=9, color=BLACK)
            ws.cell(row_idx, col_idx).alignment = Alignment(vertical="top", wrap_text=True)
    last_row = max(header_row + 1, header_row + len(records))
    add_table(ws, table_name, header_row, len(headers), last_row)
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    for col in range(1, len(headers) + 1):
        header = headers[col - 1]
        width = 45 if header in {"RAW_TEXT", "VALIDATION_MESSAGE"} else 20
        ws.column_dimensions[get_column_letter(col)].width = width


def write_config_sheet(wb: Workbook, config: dict[str, Any]) -> None:
    ws = wb.create_sheet("CONFIG_USED")
    sheet_title(ws, "Parser Configuration Used", "Configuration snapshot for audit and reproducibility.", 3)
    ws["A5"] = "SETTING"
    ws["B5"] = "VALUE"
    ws["C5"] = "TYPE"
    style_header(ws, 5, 3)
    row = 6
    for key in sorted(config.keys()):
        value = config[key]
        ws.cell(row, 1).value = key
        ws.cell(row, 2).value = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
        ws.cell(row, 3).value = type(value).__name__
        ws.cell(row, 1).font = Font(name="Aptos", size=9, color=GRAY)
        ws.cell(row, 2).font = Font(name="Aptos", size=9, color=BLUE)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 16
    ws.sheet_view.showGridLines = False


def write_output_workbook(result: ParseResult, output_path: Path, config: dict[str, Any]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_transaction_sheet(wb, result.transactions)
    write_quality_sheet(wb, result)
    write_log_sheet(wb, result.logs)
    write_balance_sheet(wb, result.balances)
    recon = build_balance_recon(result.transactions, result.balances)
    write_balance_recon_sheet(wb, recon)
    write_simple_records_sheet(
        wb,
        "UNMAPPED_ROWS",
        "Rows Requiring Parser Review",
        "Rows or PDF pages that could not be converted into a valid transaction record. Use these samples to add bank-specific aliases or adapters.",
        result.unmapped_rows,
        ["SOURCE_FILE", "SOURCE_SHEET", "SOURCE_ROW", "SOURCE_PAGE", "REASON", "RAW_TEXT"],
        "tblUnmappedRows",
    )
    duplicate_rows = [
        {
            "DUP_SOURCE_FILE": rec.get("SOURCE_FILE"),
            "DUP_SOURCE_SHEET": rec.get("SOURCE_SHEET"),
            "DUP_SOURCE_ROW": rec.get("SOURCE_ROW"),
            "DUP_ACCOUNT_NUMBER": rec.get("ACCOUNT_NUMBER"),
            "DUP_ACCOUNT_CURRENCY": rec.get("ACCOUNT_CURRENCY"),
            "DUP_ENTRY_DATE": rec.get("ENTRY_DATE"),
            "DUP_TRANSACTION_AMOUNT": rec.get("TRANSACTION_AMOUNT"),
            "DUP_BANK_REFERENCE": rec.get("BANK_REFERENCE"),
            "DUP_CUSTOMER_REFERENCE": rec.get("CUSTOMER_REFERENCE"),
            "DUP_DESCRIPTION": rec.get("TRANSACTION_DESCRIPTION"),
            "DUP_RECORD_KEY": rec.get("RECORD_KEY"),
            "DUPLICATE_OF_SOURCE_FILE": rec.get("DUPLICATE_OF_SOURCE_FILE"),
            "DUPLICATE_OF_SOURCE_ROW": rec.get("DUPLICATE_OF_SOURCE_ROW"),
        }
        for rec in result.duplicates
    ]
    write_simple_records_sheet(
        wb,
        "DUPLICATES",
        "Exact Duplicate Transactions Removed",
        "The retained output contains only one row per exact transaction key. Duplicate headers are prefixed with DUP_ so the Treasury dashboard importer reads only TRANSACTION_DETAIL.",
        duplicate_rows,
        [
            "DUP_SOURCE_FILE",
            "DUP_SOURCE_SHEET",
            "DUP_SOURCE_ROW",
            "DUP_ACCOUNT_NUMBER",
            "DUP_ACCOUNT_CURRENCY",
            "DUP_ENTRY_DATE",
            "DUP_TRANSACTION_AMOUNT",
            "DUP_BANK_REFERENCE",
            "DUP_CUSTOMER_REFERENCE",
            "DUP_DESCRIPTION",
            "DUP_RECORD_KEY",
            "DUPLICATE_OF_SOURCE_FILE",
            "DUPLICATE_OF_SOURCE_ROW",
        ],
        "tblDuplicates",
    )
    write_config_sheet(wb, config)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_summary_workbook(balances: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for col, header in enumerate(BALANCE_OUTPUT_COLUMNS, 1):
        ws.cell(1, col).value = header
        ws.cell(1, col).fill = PatternFill("solid", fgColor=DARK_BLUE)
        ws.cell(1, col).font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        ws.cell(1, col).alignment = Alignment(horizontal="center", wrap_text=True)
    for row_idx, record in enumerate(balances, 2):
        for col_idx, header in enumerate(BALANCE_OUTPUT_COLUMNS, 1):
            cell = ws.cell(row_idx, col_idx)
            cell.value = xml_safe(record.get(header))
            cell.font = Font(name="Aptos", size=9, color=GREEN)
            if header == "statement_date" and cell.value:
                cell.number_format = DATE_FORMAT
            elif header in {"opening_balance", "closing_balance"}:
                cell.number_format = AMOUNT_FORMAT
            elif header == "account_number":
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(BALANCE_OUTPUT_COLUMNS))}{max(2, len(balances) + 1)}"
    ws.sheet_view.showGridLines = False
    for col, header in enumerate(BALANCE_OUTPUT_COLUMNS, 1):
        width = 30 if header in {"source_file_name", "bank_name", "account_name", "entity_name", "mapping_note"} else 16
        ws.column_dimensions[get_column_letter(col)].width = width
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_transaction_csv(records: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=TRANSACTION_OUTPUT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for rec in records:
            row = dict(rec)
            for field in ["ENTRY_DATE", "STATEMENT_DATE", "VALUE_DATE"]:
                parsed = parse_date_value(row.get(field))
                row[field] = parsed.isoformat() if parsed else ""
            writer.writerow(row)


def validate_output(output_path: Path) -> dict[str, Any]:
    wb = load_workbook(output_path, read_only=True, data_only=False)
    try:
        expected = {"TRANSACTION_DETAIL", "QUALITY_CHECK", "PARSER_LOG", "BALANCE_SUMMARY", "BALANCE_RECON", "UNMAPPED_ROWS", "DUPLICATES", "CONFIG_USED"}
        missing = expected - set(wb.sheetnames)
        if missing:
            raise ParserError(f"Output workbook is missing sheets: {sorted(missing)}")
        ws = wb["TRANSACTION_DETAIL"]
        headers = [clean_text(ws.cell(5, c).value) for c in range(1, len(TRANSACTION_OUTPUT_COLUMNS) + 1)]
        if headers[:16] != TRANSACTION_OUTPUT_COLUMNS[:16]:
            raise ParserError("Transaction output headers do not match the Treasury dashboard schema.")
        rows = sum(1 for r in ws.iter_rows(min_row=6, max_col=1, values_only=True) if r[0] not in (None, ""))
        return {"transaction_rows": rows, "sheets": wb.sheetnames}
    finally:
        wb.close()


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Parse raw bank statements into transaction-level Excel output compatible with "
            "Treasury Payment Control Tower v6."
        )
    )
    parser.add_argument("--input", required=True, help="Raw statement file or folder. Folders are scanned recursively by default.")
    parser.add_argument("--output-dir", default="output", help="Folder for parser outputs.")
    parser.add_argument("--config", default=None, help="Optional JSON configuration file.")
    parser.add_argument("--no-recursive", action="store_true", help="Do not scan subfolders.")
    parser.add_argument("--strict", action="store_true", help="Return an error when no transaction detail is found.")
    parser.add_argument("--verbose", action="store_true", help="Print detailed parser messages.")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    input_path = Path(args.input).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_path = output_dir / "bank_transaction_detail.xlsx"
    summary_output = output_dir / "summary_rebuilt.xlsx"
    csv_output = output_dir / "bank_transaction_detail.csv"
    config_path = Path(args.config).expanduser().resolve() if args.config else None

    try:
        config = load_config(config_path)
        result = parse_all(
            input_path=input_path,
            recursive=not args.no_recursive,
            config=config,
            config_path=config_path,
            output_path=output_path,
            summary_output_path=summary_output,
        )
        write_output_workbook(result, output_path, config)
        write_summary_workbook(result.balances, summary_output)
        write_transaction_csv(result.transactions, csv_output)
        validation = validate_output(output_path)
    except Exception as exc:
        logging.exception("Parser failed") if args.verbose else logging.error("Parser failed: %s", exc)
        return 2

    print("=" * 72)
    print(f"{APP_NAME} v{APP_VERSION}")
    print(f"Input: {input_path}")
    print(f"Output workbook: {output_path}")
    print(f"Output CSV: {csv_output}")
    print(f"Balance summary: {summary_output}")
    print(f"Files scanned: {len(result.logs)}")
    print(f"Transaction rows retained: {len(result.transactions)}")
    print(f"Transaction rows requiring review: {sum(1 for rec in result.transactions if rec.get('RECORD_STATUS') != 'VALID')}")
    print(f"Exact duplicates removed: {len(result.duplicates)}")
    print(f"Balance rows retained: {len(result.balances)}")
    print(f"Unmapped/raw review rows: {len(result.unmapped_rows)}")
    print("Validation: PASS")
    print("=" * 72)

    if args.strict and not result.transactions:
        print(
            "STRICT MODE: No transaction-level rows were found. The source may contain only balance summaries, "
            "or a bank-specific adapter is required.",
            file=sys.stderr,
        )
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
